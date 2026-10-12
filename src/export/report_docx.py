"""
report_docx.py
──────────────
Word (.docx) renderer for the EIS report. It draws exactly the same blocks as
report_export.render_pdf() (see eis_report.py for the block kinds), so the PDF
and the Word file always have the same content.

Choices worth knowing:
  - Inline math ($\\tau$, $R_{ct}$) uses real Word subscript/superscript runs,
    so it stays editable text.
  - Display formulas are inserted as images (rendered LaTeX), with the LaTeX
    source in small grey text below. In Word you can retype a formula as a
    native equation: press Alt + =, paste the LaTeX source, then press Space.
  - Figures are placed side by side in a borderless table.

Requires:  pip install python-docx
Expected location: src/export/report_docx.py
"""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .latex_text import FORMULA_DPI, png_size, render_formula_png, rich_segments
from .report_export import REPORT_DIR, _ensure_dirs, safe_name

GREY = RGBColor(110, 110, 110)
PAGE_TEXT_WIDTH_CM = 17.0          # A4 (21 cm) minus 2 cm margins on each side


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _add_rich(paragraph, text: str, bold: bool = False, size: float | None = None,
              color: RGBColor | None = None, mono: bool = False) -> None:
    """Add text (with inline $...$ math) to a paragraph as formatted runs."""
    for segment, mode in rich_segments(text):
        run = paragraph.add_run(segment)
        if bold:
            run.bold = True
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if mono:
            run.font.name = "Consolas"
        if mode == "sub":
            run.font.subscript = True
        elif mode == "sup":
            run.font.superscript = True


def _add_rule(doc) -> None:
    """A thin grey horizontal line (a paragraph with a bottom border)."""
    paragraph = doc.add_paragraph()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"),
                       ("w:space", "1"), ("w:color", "B4B4B4")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _add_formula(doc, tex: str) -> None:
    png = render_formula_png(tex)
    if png is not None:
        wpx, _ = png_size(png)
        width_in = min(wpx / FORMULA_DPI, PAGE_TEXT_WIDTH_CM / 2.54)   # natural size
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))

    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = source.add_run(f"LaTeX:  {tex}")
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = GREY


def _add_figure_row(doc, items, heading: str | None, max_height_cm: float = 9.0) -> None:
    if heading:
        doc.add_heading(heading, level=1)

    n = len(items)
    col_cm = PAGE_TEXT_WIDTH_CM / n
    table = doc.add_table(rows=2, cols=n)        # row 0: pictures, row 1: captions
    for k, (caption, png) in enumerate(items):
        wpx, hpx = png_size(png)
        width = col_cm - 0.4
        if width * hpx / wpx > max_height_cm:     # tall figure (Bode): shrink it
            width = max_height_cm * wpx / hpx

        picture_cell = table.cell(0, k)
        picture_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_cell.paragraphs[0].add_run().add_picture(
            io.BytesIO(png), width=Cm(width))

        caption_par = table.cell(1, k).paragraphs[0]
        caption_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_rich(caption_par, caption, size=9, color=GREY)
    doc.add_paragraph()                           # a little air after the row


# ─────────────────────────────────────────────────────────────────────────────
# Draw the blocks
# ─────────────────────────────────────────────────────────────────────────────

def _draw(doc, blocks: list) -> None:
    for block in blocks:
        kind = block[0]
        if kind == "title":
            _add_rich(doc.add_heading("", level=0), block[1])
        elif kind == "h2":
            _add_rich(doc.add_heading("", level=1), block[1])
        elif kind == "h3":
            _add_rich(doc.add_heading("", level=2), block[1])
        elif kind == "text":
            _add_rich(doc.add_paragraph(), block[1])
        elif kind == "note":
            _add_rich(doc.add_paragraph(), block[1], size=9, color=GREY)
        elif kind == "kv":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            _add_rich(paragraph, f"{block[1]}:  ", bold=True)
            _add_rich(paragraph, block[2])
        elif kind == "bullets":
            for item in block[1]:
                _add_rich(doc.add_paragraph(style="List Bullet"), item, size=10)
        elif kind == "formula":
            _add_formula(doc, block[1])
        elif kind == "figures":
            _add_figure_row(doc, block[1], block[2] if len(block) > 2 else None)
        elif kind == "separator":
            _add_rule(doc)
        elif kind == "group":
            # "Keep with next" on every paragraph but the last one makes Word keep
            # the whole group (heading, value, formula, assumptions) on one page.
            start = len(doc.paragraphs)
            _draw(doc, block[1])
            for paragraph in doc.paragraphs[start:-1]:
                paragraph.paragraph_format.keep_with_next = True
        else:
            raise ValueError(f"Unknown block kind: {kind!r}")


def render_docx(blocks: list) -> Document:
    """Draw the blocks and return the (not yet saved) Word document."""
    doc = Document()

    section = doc.sections[0]                        # A4 with 2 cm margins
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    _draw(doc, blocks)
    return doc


def save_docx(doc: Document, name_hint: str) -> str:
    """Write the document to results/reports/<name>_<timestamp>.docx, return its path."""
    _ensure_dirs()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = REPORT_DIR / f"{safe_name(name_hint)}_{timestamp}.docx"
    doc.save(str(path))
    return str(path)
